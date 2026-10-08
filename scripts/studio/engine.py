"""Low-level python-pptx / OOXML helpers (theme-agnostic).

Colours are referenced by semantic name (looked up in C) or raw hex.
Fonts are looked up in F at call time. Both are filled by set_theme().
Lessons baked in (see references/qa.md):
  * text with letter-spacing + transparency is clipped by LibreOffice -> alpha dropped
  * negative animation offsets must be written '#ppt_y-0.05', not '#ppt_y+-0.05'
  * transitions use mc:AlternateContent with a p:fade fallback (old PowerPoint / WPS)
"""
import warnings
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
P15 = 'http://schemas.microsoft.com/office/powerpoint/2012/main'

C = {'white': 'FFFFFF', 'black': '000000'}          # semantic palette (filled by set_theme)
F = {'head': 'Noto Serif SC', 'body': 'Microsoft YaHei', 'latin': 'Palatino Linotype'}
FRAME = {'name': 'arch'}                             # default picture/card frame geometry

# frame name -> (MSO_SHAPE for autoshapes, prst for pictures, adjust values in 1/100000)
FRAMES = {
    'arch': (MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 'round2SameRect', [('adj1', 50000), ('adj2', 0)]),
    'rrect': (MSO_SHAPE.ROUNDED_RECTANGLE, 'roundRect', [('adj', 9000)]),
    'pill': (MSO_SHAPE.ROUNDED_RECTANGLE, 'roundRect', [('adj', 50000)]),
    'plaque': (MSO_SHAPE.PLAQUE, 'plaque', [('adj', 14000)]),
    'oval': (MSO_SHAPE.OVAL, 'ellipse', []),
    'snip': (MSO_SHAPE.SNIP_2_SAME_RECTANGLE, 'snip2SameRect', [('adj1', 16000), ('adj2', 0)]),
    'octagon': (MSO_SHAPE.OCTAGON, 'octagon', [('adj', 22000)]),
    'rect': (MSO_SHAPE.RECTANGLE, 'rect', []),
    'diamond': (MSO_SHAPE.DIAMOND, 'diamond', []),
    'tri': (MSO_SHAPE.ISOSCELES_TRIANGLE, 'triangle', []),
}


def set_theme(theme):
    C.clear(); C.update({'white': 'FFFFFF', 'black': '000000'}); C.update(theme['colors'])
    F.update(theme.get('fonts', {}))
    FRAME['name'] = theme.get('frame', 'arch')


def frame_key(kind):
    return FRAME['name'] if kind in (None, 'frame') else kind


def hexof(c):
    return C.get(c, c)


def rgb(c):
    return RGBColor.from_string(hexof(c))


def E(v):
    return int(round(v * 914400))


def _set_av(prstGeom, pairs):
    av = prstGeom.find(qn('a:avLst'))
    if av is None:
        av = etree.SubElement(prstGeom, qn('a:avLst'))
    for g in list(av):
        av.remove(g)
    for n, v in pairs:
        g = etree.SubElement(av, qn('a:gd')); g.set('name', n); g.set('fmla', f'val {v}')


# ---------------------------------------------------------------- text
def set_fonts(rPr, latin, ea):
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
    etree.SubElement(rPr, qn('a:latin')).set('typeface', latin)
    etree.SubElement(rPr, qn('a:ea')).set('typeface', ea)
    etree.SubElement(rPr, qn('a:cs')).set('typeface', latin)


def style_run(run, size=14, color='ink', bold=False, italic=False, font=None, latin=None,
              spacing=None, alpha=None):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    if alpha is not None and spacing:
        warnings.warn('letter-spacing + transparency gets clipped in LibreOffice; alpha dropped')
        alpha = None
    if alpha is not None:
        srgb = rPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        etree.SubElement(srgb, qn('a:alpha')).set('val', str(int(alpha * 1000)))
    if spacing:
        rPr.set('spc', str(int(spacing * 100)))
    font = F.get(font, font) if font else F['body']          # 'head' / 'body' / 'latin' aliases allowed
    latin = F.get(latin, latin) if latin else None
    set_fonts(rPr, latin or font, font)


def add_text(slide, x, y, w, h, paras, size=14, color='ink', bold=False, font=None, latin=None,
             align='l', anchor='t', line=None, spacing=None, italic=False, para_after=0, name=None,
             wrap=True, alpha=None, rot=0, vert=False):
    """paras: str | list of paragraphs; a paragraph is str | dict | list of run dicts.
    run dict keys: text, size, color, bold, italic, font, latin, spacing, alpha.
    Use latin=True inside a run to mean 'the theme latin font'."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    if name:
        tb.name = name
    if rot:
        tb.rotation = rot
    tf = tb.text_frame
    tf.word_wrap = wrap
    bp = tf._txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bp.set(k, '0')
    if vert:
        bp.set('vert', 'vert')   # real vertical text; rotating a box whose frame leaves the slide gets clipped
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, (str, dict)):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT, 'j': PP_ALIGN.JUSTIFY}[align]
        if line:
            para.line_spacing = line
        if para_after:
            para.space_after = Pt(para_after)
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            if isinstance(r, str):
                r = dict(text=r)
            lat = r.get('latin', latin)
            if lat is True:
                lat = F['latin']
            run = para.add_run()
            run.text = r['text']
            style_run(run, size=r.get('size', size), color=r.get('color', color), bold=r.get('bold', bold),
                      italic=r.get('italic', italic), font=r.get('font', font), latin=lat,
                      spacing=r.get('spacing', spacing), alpha=r.get('alpha', alpha))
    return tb


# ---------------------------------------------------------------- shapes
def _spPr(shape):
    return shape._element.spPr


def set_fill(shape, color, alpha=100):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if alpha < 100:
        srgb = _spPr(shape).find(qn('a:solidFill')).find(qn('a:srgbClr'))
        etree.SubElement(srgb, qn('a:alpha')).set('val', str(int(alpha * 1000)))


def set_line(shape, color=None, width=1.0, alpha=100, dash=None):
    if color is None:
        shape.line.fill.background()
        return
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    ln = _spPr(shape).find(qn('a:ln'))
    if alpha < 100:
        srgb = ln.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        etree.SubElement(srgb, qn('a:alpha')).set('val', str(int(alpha * 1000)))
    if dash:
        etree.SubElement(ln, qn('a:prstDash')).set('val', dash)


def _strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def add_shape(slide, kind, x, y, w, h, fill=None, alpha=100, line=None, lw=1.0, lalpha=100,
              name=None, adj=None, dash=None, rot=0):
    """kind: a FRAMES key ('arch','rrect','pill','plaque','oval','snip','octagon','rect','diamond','tri')
    or 'frame' (= theme frame)."""
    k = frame_key(kind)
    mso, prst, pairs = FRAMES[k]
    sh = slide.shapes.add_shape(mso, E(x), E(y), E(w), E(h))
    if name:
        sh.name = name
    pg = _spPr(sh).find(qn('a:prstGeom'))
    if adj is not None:
        pairs = [(n, int(v * 100000)) for (n, _), v in zip(pairs, adj)]
    if pairs:
        _set_av(pg, pairs)
    if fill:
        set_fill(sh, fill, alpha)
    else:
        sh.fill.background()
    set_line(sh, line, lw, lalpha, dash)
    if rot:
        sh.rotation = rot
    _strip_style(sh)
    sh.text_frame.text = ''
    return sh


def add_line(slide, x1, y1, x2, y2, color='line', width=0.75, alpha=100, dash=None, name=None):
    ln = slide.shapes.add_connector(1, E(x1), E(y1), E(x2), E(y2))
    if name:
        ln.name = name
    set_line(ln, color, width, alpha, dash)
    _strip_style(ln)
    return ln


def freeform(slide, pts, color, width=2.25, dash=None, arrow=True, name=None):
    fb = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(px), E(py)) for px, py in pts[1:]], close=False)
    sh = fb.convert_to_shape()
    if name:
        sh.name = name
    sh.fill.background()
    set_line(sh, color, width, dash=dash)
    ln = _spPr(sh).find(qn('a:ln'))
    etree.SubElement(ln, qn('a:round'))
    if arrow:
        te = etree.SubElement(ln, qn('a:tailEnd')); te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    _strip_style(sh)
    return sh


def shadow(shape, blur=0.18, dist=0.06, dirdeg=90, color='3B2E25', alpha=28):
    spPr = _spPr(shape)
    old = spPr.find(qn('a:effectLst'))
    if old is not None:
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'))
    sh.set('blurRad', str(E(blur))); sh.set('dist', str(E(dist)))
    sh.set('dir', str(int(dirdeg * 60000))); sh.set('algn', 'ctr'); sh.set('rotWithShape', '0')
    c = etree.SubElement(sh, qn('a:srgbClr')); c.set('val', hexof(color))
    etree.SubElement(c, qn('a:alpha')).set('val', str(alpha * 1000))
    ln = spPr.find(qn('a:ln'))
    if ln is not None:
        spPr.remove(eff); ln.addnext(eff)


# ---------------------------------------------------------------- pictures
def add_pic(slide, path, x, y, w=None, h=None, cover=True, focus=(0.5, 0.5), geom=None, name=None, alpha=None):
    """cover=True crops to fill w*h around focus (0..1, 0..1). geom: FRAMES key / 'frame' / prst name."""
    iw, ih = Image.open(path).size
    if w is None:
        w = h * iw / ih
    if h is None:
        h = w * ih / iw
    pic = slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    if name:
        pic.name = name
    if cover:
        ar_box, ar_img = w / h, iw / ih
        if ar_img > ar_box:
            keep = ar_box / ar_img; off = (1 - keep) * focus[0]
            pic.crop_left, pic.crop_right = off, 1 - keep - off
        elif ar_img < ar_box:
            keep = ar_img / ar_box; off = (1 - keep) * focus[1]
            pic.crop_top, pic.crop_bottom = off, 1 - keep - off
    if geom:
        pg = pic._element.spPr.find(qn('a:prstGeom'))
        k = frame_key(geom)
        if k in FRAMES:
            _, prst, pairs = FRAMES[k]
            pg.set('prst', prst)
            if pairs:
                _set_av(pg, pairs)
        else:
            pg.set('prst', geom)
    if alpha is not None:
        blip = pic._element.find('.//' + qn('a:blip'))
        etree.SubElement(blip, qn('a:alphaModFix')).set('amt', str(int(alpha * 1000)))
    return pic


def pic_outline(pic, color='line', width=0.75):
    spPr = pic._element.spPr
    ln = etree.SubElement(spPr, qn('a:ln')); ln.set('w', str(int(width * 12700)))
    sf = etree.SubElement(ln, qn('a:solidFill')); etree.SubElement(sf, qn('a:srgbClr')).set('val', hexof(color))


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- animation
class Timeline:
    """Auto-playing entrance timeline (starts when the slide appears, like the reference deck).
    add(shape, effect, at=ms, dur=ms, **kw). Effects: fade wipe(dir) rise(dy) slide(dx)
    zoom(scale) circle split splitv letters(pct)."""

    def __init__(self):
        self.items = []
        self.t = 0

    def add(self, shape, effect='fade', at=None, dur=500, **kw):
        if shape is None:
            return
        if at is None:
            at = self.t
        self.items.append(dict(shape=shape, effect=effect, at=int(at), dur=int(dur), kw=kw))
        self.t = max(self.t, at)

    def end(self):
        return max([i['at'] + i['dur'] for i in self.items] or [0])


def _tgt(spid):
    return f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'


def _set_vis(nid, spid):
    return (f'<p:set><p:cBhvr><p:cTn id="{nid}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'{_tgt(spid)}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:to><p:strVal val="visible"/></p:to></p:set>')


def _fx(nid, spid, filt, dur):
    return (f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{nid}" dur="{dur}"/>'
            f'{_tgt(spid)}</p:cBhvr></p:animEffect>')


def _off(base, d):
    return f'{base}+{d}' if d >= 0 else f'{base}-{abs(d)}'


def _anim(nid, spid, attr, v0, v1, dur):
    val = lambda v: f'<p:fltVal val="{v}"/>' if isinstance(v, (int, float)) else f'<p:strVal val="{v}"/>'
    return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{nid}" dur="{dur}" fill="hold"/>'
            f'{_tgt(spid)}<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:tavLst><p:tav tm="0"><p:val>{val(v0)}</p:val></p:tav><p:tav tm="100000"><p:val>{val(v1)}</p:val></p:tav></p:tavLst></p:anim>')


WIPE = dict(left=(8, 'wipe(left)'), right=(2, 'wipe(right)'), up=(4, 'wipe(down)'), down=(1, 'wipe(up)'))


def _effect_xml(ids, item, node_type, grp):
    spid = item['shape'].shape_id
    eff, dur, kw = item['effect'], item['dur'], item['kw']
    cid = next(ids)
    body = _set_vis(next(ids), spid)
    extra = ''
    if eff == 'fade':
        pid, sub = 10, 0; body += _fx(next(ids), spid, 'fade', dur)
    elif eff == 'wipe':
        sub, filt = WIPE[kw.get('dir', 'left')]; pid = 22; body += _fx(next(ids), spid, filt, dur)
    elif eff == 'rise':
        pid, sub = 42, 0
        body += _fx(next(ids), spid, 'fade', dur)
        body += _anim(next(ids), spid, 'ppt_x', '#ppt_x', '#ppt_x', dur)
        body += _anim(next(ids), spid, 'ppt_y', _off('#ppt_y', kw.get('dy', 0.1)), '#ppt_y', dur)
    elif eff == 'slide':
        pid, sub = 42, 0
        body += _fx(next(ids), spid, 'fade', dur)
        body += _anim(next(ids), spid, 'ppt_x', _off('#ppt_x', kw.get('dx', 0.05)), '#ppt_x', dur)
        body += _anim(next(ids), spid, 'ppt_y', '#ppt_y', '#ppt_y', dur)
    elif eff == 'zoom':
        pid, sub = 53, 16
        s0 = kw.get('scale', 0)
        body += _anim(next(ids), spid, 'ppt_w', s0 if s0 == 0 else f'#ppt_w*{s0}', '#ppt_w', dur)
        body += _anim(next(ids), spid, 'ppt_h', s0 if s0 == 0 else f'#ppt_h*{s0}', '#ppt_h', dur)
        body += _fx(next(ids), spid, 'fade', dur)
    elif eff == 'circle':
        pid, sub = 6, 32; body += _fx(next(ids), spid, 'circle(out)', dur)
    elif eff == 'split':
        pid, sub = 16, 42; body += _fx(next(ids), spid, 'barn(outHorizontal)', dur)
    elif eff == 'splitv':
        pid, sub = 16, 26; body += _fx(next(ids), spid, 'barn(outVertical)', dur)
    elif eff == 'letters':
        pid, sub = 0, 0
        body += _fx(next(ids), spid, 'fade', dur)
        body += _anim(next(ids), spid, 'ppt_x', '#ppt_x+0.03', '#ppt_x', dur)
        extra = f'<p:iterate type="lt"><p:tmPct val="{kw.get("pct", 12000)}"/></p:iterate>'
    else:
        raise ValueError(eff)
    decel = ' decel="100000"' if eff in ('rise', 'slide', 'letters', 'zoom') else ''
    gid = f' grpId="{grp}"' if grp is not None else ''
    return (f'<p:par><p:cTn id="{cid}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold"{gid}{decel} nodeType="{node_type}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst>{extra}<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')


def build_timing(slide, tl):
    if not tl.items:
        return
    ids = iter(range(4, 100000))
    steps, bld, first = '', {}, True
    for st in sorted(set(i['at'] for i in tl.items)):
        inner = ''
        for k, it in enumerate([i for i in tl.items if i['at'] == st]):
            is_sp = it['shape']._element.tag == qn('p:sp')
            inner += _effect_xml(ids, it, 'afterEffect' if (k == 0 and not first) else 'withEffect', 0 if is_sp else None)
            if is_sp:
                bld[it['shape'].shape_id] = bool(it['shape'].has_text_frame and it['shape'].text_frame.text.strip())
        steps += (f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="{st}"/></p:stCondLst>'
                  f'<p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>')
        first = False
    bld_xml = ''.join(f'<p:bldP spid="{s}" grpId="0"{"" if t else " animBg=\"1\""}/>' for s, t in bld.items())
    xml = (f'<p:timing xmlns:p="{NS_P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst>'
           f'<p:childTnLst>{steps}</p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst>' + (f'<p:bldLst>{bld_xml}</p:bldLst>' if bld_xml else '') + '</p:timing>')
    sld = slide._element
    old = sld.find(qn('p:timing'))
    if old is not None:
        sld.remove(old)
    el = etree.fromstring(xml)
    ext = sld.find(qn('p:extLst'))
    ext.addprevious(el) if ext is not None else sld.append(el)


def set_transition(slide, kind='fade', dur=1000, **kw):
    """kind: 'fade' | 'p14:<name>' (vortex, gallery, reveal, prism, ripple, flip...) |
    'p15:<prst>' (curtains, pageCurlDouble, prestige, drape, wind, airplane, origami, fracture...)."""
    sld = slide._element
    for old in sld.findall(qn('p:transition')) + sld.findall('{%s}AlternateContent' % MC):
        sld.remove(old)
    spd = 'slow' if dur >= 1000 else 'med'
    if kind == 'fade':
        el = etree.fromstring(f'<p:transition xmlns:p="{NS_P}" spd="{spd}"><p:fade/></p:transition>')
    else:
        ns, name = kind.split(':')
        attrs = ''.join(f' {k}="{v}"' for k, v in kw.items())
        inner = f'<p14:{name}{attrs}/>' if ns == 'p14' else f'<p15:prstTrans prst="{name}"{attrs}/>'
        el = etree.fromstring(
            f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{NS_P}" xmlns:p14="{P14}" xmlns:p15="{P15}">'
            f'<mc:Choice Requires="{ns}"><p:transition spd="slow" p14:dur="{dur}">{inner}</p:transition></mc:Choice>'
            f'<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    sld.find(qn('p:clrMapOvr')).addnext(el)


# ---------------------------------------------------------------- package-level
def set_master_background(prs, image_path):
    master = prs.slide_master
    _, rId = master.part.get_or_add_image_part(image_path)
    bg = etree.fromstring(
        f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:bgPr><a:blipFill dpi="0" rotWithShape="1">'
        f'<a:blip r:embed="{rId}"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    cSld = master._element.find(qn('p:cSld'))
    for old in cSld.findall(qn('p:bg')):      # the default template already has one -> two bg = ignored
        cSld.remove(old)
    cSld.insert(0, bg)


def set_master_color(prs, color):
    master = prs.slide_master
    bg = etree.fromstring(f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}"><p:bgPr><a:solidFill><a:srgbClr val="{hexof(color)}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>')
    cSld = master._element.find(qn('p:cSld'))
    for old in cSld.findall(qn('p:bg')):
        cSld.remove(old)
    cSld.insert(0, bg)


def apply_theme_xml(prs, theme):
    """Write palette + fonts into the theme part, so new text / charts follow the style."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    tp = prs.slide_master.part.part_related_by(RT.THEME)
    th = etree.fromstring(tp.blob)
    na = {'a': NS_A}
    th.set('name', theme.get('name', 'Studio'))
    c = theme['colors']
    pal = dict(dk1=c['ink'], lt1='FFFFFF', dk2=c.get('primary_dk', c['primary']), lt2=c['paper'],
               accent1=c['primary'], accent2=c['accent'], accent3=c['line'], accent4=c.get('primary_lt', c['primary']),
               accent5=c.get('extra', c['accent']), accent6=c.get('accent2', c['accent']),
               hlink=c.get('deep', c['ink']), folHlink=c['ink_soft'])
    cs = th.find('.//a:clrScheme', na); cs.set('name', theme.get('name', 'Studio'))
    for k, v in pal.items():
        el = cs.find('a:' + k, na)
        for ch in list(el):
            el.remove(ch)
        etree.SubElement(el, qn('a:srgbClr')).set('val', v)
    fs = th.find('.//a:fontScheme', na); fs.set('name', theme.get('name', 'Studio'))
    for tag, lat, ea in (('a:majorFont', F['latin'], F['head']), ('a:minorFont', F['body'], F['body'])):
        f = fs.find(tag, na)
        f.find('a:latin', na).set('typeface', lat)
        f.find('a:ea', na).set('typeface', ea)
    tp._blob = etree.tostring(th, xml_declaration=True, encoding='UTF-8', standalone=True)


def drop_unused_layouts(prs, keep):
    for lay in list(prs.slide_master.slide_layouts):
        if lay is not keep and not lay.used_by_slides:
            prs.slide_master.slide_layouts.remove(lay)
