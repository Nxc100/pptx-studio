"""Deck = theme + assets + slide archetypes.  Usage: see examples/ and references/archetypes.md

    from studio import Deck
    d = Deck('themes/medieval_pastoral.json', assets='assets', brand='Ba · 无界森居', brand_en='MEDIEVAL PASTORAL INTERIOR')
    d.cover(...); d.contents(...); d.section(...); ...; d.save('out.pptx')

Every archetype animates itself (auto-play timeline) and gets the theme transition for its kind.
Colours are semantic names from the theme: paper ink ink_soft primary primary_dk primary_lt accent
accent2 extra line line_lt card muted deep white (or raw hex).
"""
import os, json
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.util import Pt
from .engine import *
from .engine import set_fonts

W, H = 13.333, 7.5
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))          # pptx-studio/


def _load_theme(theme):
    """theme: dict | path to json | name of themes/<name>.json"""
    if isinstance(theme, dict):
        return theme
    p = theme if os.path.exists(theme) else os.path.join(ROOT, 'themes', theme if theme.endswith('.json') else theme + '.json')
    T = json.load(open(p, encoding='utf8'))
    T.setdefault('_dir', os.path.join(os.path.dirname(os.path.abspath(p)), T.get('asset_dir', '')))
    return T


class Deck:
    def __init__(self, theme, assets='.', brand='', brand_en='', meta=None):
        self.T = _load_theme(theme)
        set_theme(self.T)
        # theme-owned decorations/icons live in themes/<asset_dir>/ ; project photos/renders live in `assets`
        self.theme_dir = self.T.get('_dir') or os.path.join(ROOT, 'themes', self.T.get('asset_dir', ''))
        self.assets = os.path.abspath(assets)
        self.brand, self.brand_en = brand, brand_en
        self.meta = meta or {}
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = E(W), E(H)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0
        bg = self._deco_path('paper')
        if bg and os.path.exists(bg):
            set_master_background(self.prs, bg)
        else:
            set_master_color(self.prs, 'paper')

    # ------------------------------------------------------------ assets
    def A(self, rel):
        return rel if os.path.isabs(rel) else os.path.join(self.assets, rel)

    def has(self, rel):
        return bool(rel) and os.path.exists(self.A(rel))

    def _deco_path(self, key):
        """theme deco key -> absolute path (theme dir); otherwise treat key as a project asset path"""
        d = self.T.get('deco', {})
        if key in d and isinstance(d[key], str):
            return d[key] if os.path.isabs(d[key]) else os.path.join(self.theme_dir, d[key])
        return self.A(key)

    def deco(self, s, key, x, y, w=None, h=None, alpha=None):
        """key: theme deco key (corner_tl branch_tr sprig_br garland ornament bullet node divider seal watermark)
        or an asset path. Missing assets are silently skipped (returns None) so themes may omit any of them."""
        p = self._deco_path(key)
        if not p or not os.path.exists(p):
            return None
        return add_pic(s, p, x, y, w, h, cover=False, alpha=alpha)

    def floaters(self):
        return self.T.get('deco', {}).get('floaters', [])

    def icon(self, s, name, x, y, size=0.42, variant='primary'):
        """icons/<name>_<suffix>.png in the theme dir; variant: primary | light | accent (mapped by theme['icons'])"""
        ic = self.T.get('icons', {})
        p = os.path.join(self.theme_dir, ic.get('dir', 'icons'), f"{name}_{ic.get(variant, variant)}.png")
        if not os.path.exists(p):
            return None
        return add_pic(s, p, x, y, size, size, cover=False)

    # ------------------------------------------------------------ slide plumbing
    def new(self, kind='content'):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        tr = self.T.get('transitions', {}).get(kind) or self.T.get('transitions', {}).get('content') or ['fade', 800]
        set_transition(s, tr[0], tr[1], **(tr[2] if len(tr) > 2 else {}))
        return s, Timeline(), self.n

    def finish(self, s, tl, note=''):
        build_timing(s, tl)
        if not note:   # auto speaker notes: title + the slide's main texts (spec can always pass a real note)
            skip = {self.brand, self.brand_en}
            texts = [sh.text_frame.text.replace(chr(10), ' ').strip() for sh in s.shapes if sh.has_text_frame]
            texts = [t for t in texts if len(t) >= 4 and t not in skip and not t.startswith('—') and self.brand not in t[:len(self.brand) + 1]]
            note = '；'.join(texts)[:220]
        notes(s, note)
        return s

    def save(self, path):
        apply_theme_xml(self.prs, self.T)
        drop_unused_layouts(self.prs, self.blank)
        cp = self.prs.core_properties
        for k in ('title', 'author', 'subject', 'keywords'):
            if self.meta.get(k):
                setattr(cp, k, self.meta[k])
        if self.meta.get('author'):
            cp.last_modified_by = self.meta['author']
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.prs.save(path)
        return path

    # ------------------------------------------------------------ components
    def chrome(self, s, n, tl, title, en, sprig=True, header=True):
        """content-slide frame: corner sprig, title, english caption, running header, folio, optional bottom sprig"""
        txt = self.T.get('text', {})
        sp = self.deco(s, 'corner_tl', -0.55, -0.4, w=2.1)
        t = add_text(s, 1.45, 0.34, 7.5, 0.62, title, size=26, color='primary_dk', font=F['head'], bold=True)
        e = add_text(s, 1.47, 0.98, 7.5, 0.28, en.upper(), size=9.5, color='line', latin=F['latin'], italic=True, spacing=2.6)
        tl.add(sp, 'fade', at=0, dur=700)
        tl.add(t, 'wipe', at=150, dur=550)
        tl.add(e, 'wipe', at=450, dur=550)
        if header and (self.brand or self.brand_en):
            hd = add_text(s, 9.0, 0.38, 3.75, 0.55,
                          [[dict(text=self.brand, font=F['head'], size=11, color='ink_soft')],
                           [dict(text=self.brand_en, size=7.5, color='line', latin=F['latin'], spacing=2)]], align='r')
            tl.add(hd, 'fade', at=350, dur=600)
        fl = add_text(s, 11.95, 6.98, 0.8, 0.3, txt.get('folio', '— {n:02d} —').format(n=n), size=9.5, color='line',
                      latin=F['latin'], align='r')
        if sprig:
            lv = self.deco(s, 'sprig_br', 12.35, 5.55, h=2.25)
            tl.add(lv, 'rise', at=500, dur=900, dy=0.05)
        tl.t = 700
        return t

    def framed_photo(self, s, path, x, y, w, h, frame=0.13, focus=(0.5, 0.5), shadow_on=True, double=False, geom='frame'):
        frames = []
        if frame:
            frames.append(add_shape(s, geom, x - frame, y - frame, w + 2 * frame, h + 2 * frame, line='line', lw=0.9))
            if double:
                frames.append(add_shape(s, geom, x - frame - 0.08, y - frame - 0.08, w + 2 * frame + 0.16,
                                        h + 2 * frame + 0.16, line='line', lw=0.5, lalpha=70))
        p = add_pic(s, self.A(path), x, y, w, h, geom=geom, focus=focus)
        if shadow_on:
            shadow(p, blur=0.28, dist=0.07, alpha=30)
        return p, frames

    def card(self, s, x, y, w, h, fill='card', alpha=70, line='line', lw=0.6, kind='rect', shadow_on=False):
        c = add_shape(s, kind, x, y, w, h, fill=fill, alpha=alpha, line=line, lw=lw)
        if shadow_on:
            shadow(c, blur=0.2, dist=0.05, alpha=18)
        return c

    def icon_badge(self, s, name, x, y, d=0.62, fill='primary', variant='light', ring=True):
        items = [add_shape(s, 'oval', x, y, d, d, fill=fill)]
        if ring:
            items.append(add_shape(s, 'oval', x - 0.06, y - 0.06, d + 0.12, d + 0.12, line='line', lw=0.6))
        ic = self.icon(s, name, x + d * 0.22, y + d * 0.22, d * 0.56, variant)
        if ic is not None:
            items.append(ic)
        return items

    def side_heading(self, s, x, y, cn, en, tl, at):
        t = add_text(s, x, y, 4.2, 0.36, [[dict(text=cn + '　', bold=True),
                                           dict(text=en.upper(), latin=True, italic=True, size=8.5, color='line', bold=False)]],
                     size=14, color='primary_dk', font=F['head'])
        tl.add(t, 'wipe', at=at, dur=550)
        return t

    def fit(self, rel, box):
        iw, ih = Image.open(self.A(rel)).size
        sc = min(box[2] / iw, box[3] / ih)
        pw, ph = iw * sc, ih * sc
        return box[0] + (box[2] - pw) / 2, box[1] + (box[3] - ph) / 2, pw, ph, sc

    @staticmethod
    def pmap(px, py, box, img_size, offset=(0, 0)):
        x, y, w, h = box
        return x + (px - offset[0]) * w / img_size[0], y + (py - offset[1]) * h / img_size[1]

    def points_list(self, s, x, y, w, pts, tl, at, gap=0.86, size=10.5):
        fills = ['primary', 'accent', 'line', 'accent2']
        for i, (ic, k, v) in enumerate(pts):
            yy = y + i * gap
            bd = self.icon_badge(s, ic, x, yy + 0.02, d=0.42, ring=False, fill=fills[i % 4])
            t = add_text(s, x + 0.58, yy - 0.03, w - 0.58, gap, [[dict(text=k, font=F['head'], bold=True, color='ink', size=12)], v],
                         size=size, color='ink_soft', font=F['body'], line=1.28)
            for b in bd:
                tl.add(b, 'zoom', at=at + i * 200, dur=400)
            tl.add(t, 'slide', at=at + 80 + i * 200, dur=600, dx=0.03)

    def bullets(self, s, x, y, w, items, tl, at, gap=0.42, size=13, key_color='primary_dk'):
        """items: [(key, value)] rendered as 'key  value' rows with theme bullet ornament"""
        for i, (k, v) in enumerate(items):
            yy = y + i * gap
            b = self.deco(s, 'bullet', x, yy + 0.07, w=0.18)
            t = add_text(s, x + 0.31, yy, w - 0.31, gap, [[dict(text=k + '　', color=key_color, font=F['head'], bold=True), dict(text=v, color='ink')]],
                         size=size, font=F['body'], line=1.2)
            tl.add(b, 'zoom', at=at + i * 150, dur=400)
            tl.add(t, 'slide', at=at + i * 150, dur=600, dx=0.03)

    def _pill(self, s, x, y, w, h, text, fill='primary', size=9, tl=None, at=0):
        p = add_shape(s, 'pill', x, y, w, h, fill=fill)
        t = add_text(s, x, y, w, h, text, size=size, color='white', latin=F['latin'], align='c', anchor='m', spacing=2)
        if tl:
            tl.add(p, 'wipe', at=at, dur=500); tl.add(t, 'wipe', at=at, dur=500)
        return p, t

    # ============================================================ archetypes
    def cover(self, title, subtitle, en_line, info, photo, focus=(0.5, 0.5), kicker='INTERIOR DESIGN PROPOSAL', note=''):
        """title: str or list of run dicts (e.g. [dict(text='Ba', latin=True, italic=True), dict(text=' · 无界森居')])
        info: [(label, value)] e.g. [('主讲人','刘诗云'), ('学　号','2024…'), ('项目地点','…')]"""
        s, tl, n = self.new('cover')
        q = self.deco(s, 'ornament', -1.15, 4.55, w=3.7, alpha=30)
        aw = self.deco(s, 'watermark', 5.55, 3.2, h=3.9, alpha=22)
        p, fr = self.framed_photo(s, photo, 7.45, 0.95, 4.7, 6.0, frame=0.17, focus=focus, double=True)
        ob = self.deco(s, 'branch_tr', 9.35, -0.42, w=4.45)
        rv = self.deco(s, 'garland', 5.65, 5.72, w=3.3)
        top = add_text(s, 1.02, 0.78, 5.6, 0.3, kicker, size=10, color='line', latin=F['latin'], spacing=3)
        runs = [dict(text=title)] if isinstance(title, str) else title
        t1 = add_text(s, 0.95, 1.38, 6.6, 1.15, [runs], size=60 if len(''.join(r['text'] for r in runs)) <= 10 else 46,
                      color='primary_dk', font=F['head'], bold=True)
        seal = self.deco(s, 'seal', 5.95, 0.62, w=1.0)
        sub = add_text(s, 1.02, 2.72, 6.3, 0.5, subtitle, size=19, color='ink', font=F['head'])
        enl = add_text(s, 1.03, 3.3, 6.3, 0.3, en_line, size=8.5, color='ink_soft', latin=F['latin'], italic=True, spacing=1.2)
        dv = self.deco(s, 'divider', 1.0, 3.78, w=4.9)
        self._pill(s, 1.02, 4.32, 2.0, 0.38, self.T.get('text', {}).get('pill', 'INTERIOR  DESIGN'), tl=tl, at=2050)
        rows = []
        for i, (k, v) in enumerate(info):
            y = 5.0 + i * 0.48
            b = self.deco(s, 'bullet', 1.03, y + 0.06, w=0.2)
            t = add_text(s, 1.35, y, 5.2, 0.34, [[dict(text=k, color='line', font=F['head'], bold=True),
                                                   dict(text='　|　', color='line_lt'), dict(text=v, color='ink')]], size=13.5, font=F['head'])
            rows.append((b, t))
        tl.add(q, 'fade', at=0, dur=1200); tl.add(aw, 'fade', at=200, dur=1200)
        tl.add(p, 'zoom', at=0, dur=1300, scale=0.9)
        for i, f in enumerate(fr):
            tl.add(f, 'fade', at=500 + i * 200, dur=900)
        tl.add(top, 'wipe', at=300, dur=600)
        tl.add(t1, 'letters', at=500, dur=500, pct=14000)
        tl.add(seal, 'zoom', at=1500, dur=450)
        tl.add(sub, 'fade', at=1350, dur=700); tl.add(enl, 'wipe', at=1600, dur=700); tl.add(dv, 'split', at=1850, dur=600)
        for i, (b, t) in enumerate(rows):
            tl.add(b, 'zoom', at=2300 + i * 180, dur=400); tl.add(t, 'slide', at=2300 + i * 180, dur=600, dx=0.03)
        tl.add(ob, 'rise', at=900, dur=1200, dy=-0.06); tl.add(rv, 'rise', at=1200, dur=1200, dy=0.06)
        return self.finish(s, tl, note)

    def contents(self, parts, photo, focus=(0.5, 0.5), heading='目　录', heading_en='CONTENTS', note=''):
        """parts: [(num, cn, en, desc)] — up to 4 rows"""
        s, tl, n = self.new('contents')
        ob = self.deco(s, 'branch_tr', 10.15, -0.4, w=3.45)
        t = add_text(s, 1.0, 0.5, 4.0, 0.9, heading, size=40, color='primary_dk', font=F['head'], bold=True)
        te = add_text(s, 1.05, 1.36, 4.0, 0.3, heading_en, size=12, color='line', latin=F['latin'], italic=True, spacing=4)
        p, fr = self.framed_photo(s, photo, 1.05, 2.05, 3.55, 4.75, frame=0.15, focus=focus, double=True)
        lv = self.deco(s, 'sprig_br', 0.05, 5.35, h=2.2)
        ln = add_line(s, 5.85, 1.0, 5.85, 6.85, color='line', width=0.9)
        tl.add(ob, 'rise', at=600, dur=1000, dy=-0.05); tl.add(t, 'wipe', at=0, dur=600); tl.add(te, 'wipe', at=300, dur=600)
        tl.add(p, 'rise', at=200, dur=1000, dy=0.06)
        for i, f in enumerate(fr):
            tl.add(f, 'fade', at=600 + i * 150, dur=800)
        tl.add(lv, 'rise', at=900, dur=900, dy=0.05); tl.add(ln, 'wipe', at=700, dur=900, dir='down')
        step = 5.68 / max(len(parts), 4)
        for i, (num, cn, ens, desc) in enumerate(parts):
            y = 1.3 + i * step
            node = self.deco(s, 'node', 5.5, y + 0.02, w=0.7) or add_shape(s, 'oval', 5.52, y + 0.04, 0.66, 0.66, fill='primary')
            nt = add_text(s, 5.5, y + 0.02, 0.7, 0.7, num, size=13, color='white', latin=F['latin'], align='c', anchor='m', bold=True)
            tt = add_text(s, 6.65, y - 0.06, 5.5, 0.55, cn, size=24, color='ink', font=F['head'], bold=True)
            et = add_text(s, 6.67, y + 0.52, 5.5, 0.26, ens.upper(), size=9, color='line', latin=F['latin'], italic=True, spacing=2.2)
            dt = add_text(s, 6.67, y + 0.83, 5.8, 0.3, desc, size=11, color='ink_soft', font=F['body'])
            at = 1100 + i * 380
            tl.add(node, 'zoom', at=at, dur=450); tl.add(nt, 'zoom', at=at, dur=450)
            tl.add(tt, 'slide', at=at + 120, dur=650, dx=0.04); tl.add(et, 'wipe', at=at + 300, dur=500); tl.add(dt, 'fade', at=at + 400, dur=600)
        return self.finish(s, tl, note)

    def section(self, num, part, cn, en, photo, focus=(0.5, 0.5), items='', poem='', mark='', note=''):
        """chapter divider. part: 'One'/'Two'… ; items: sub-topics (use \\n for 2 lines); mark: vertical spine word"""
        s, tl, n = self.new('section')
        p, fr = self.framed_photo(s, photo, 7.55, 0.95, 4.95, 6.75, frame=0.17, focus=focus, double=True)
        wm = add_text(s, 12.72, 2.3, 0.5, 4.6, mark, size=26, color='line_lt', latin=F['latin'], spacing=9, vert=True,
                      align='c', anchor='m') if mark else None
        ob = self.deco(s, 'branch_tr', 10.2, -0.45, w=3.3)
        q = self.deco(s, 'ornament', -0.8, 1.15, w=3.2, alpha=28)
        big = add_text(s, 0.75, 1.0, 5.5, 2.4, num, size=150, color='muted', latin=F['latin'], italic=True, alpha=85)
        pt = add_text(s, 1.12, 1.05, 4.5, 0.45, f"{self.T.get('text', {}).get('part_word', 'The Part')} {part}", size=22,
                      color='line', latin=F['latin'], italic=True, spacing=2)
        t = add_text(s, 1.08, 2.62, 6.2, 1.1, cn, size=54 if len(cn) <= 5 else 44, color='primary_dk', font=F['head'], bold=True)
        et = add_text(s, 1.14, 3.78, 6.0, 0.3, en.upper(), size=12, color='line', latin=F['latin'], spacing=4)
        self._pill(s, 1.14, 4.32, 1.95, 0.36, self.T.get('text', {}).get('pill', 'INTERIOR  DESIGN'), size=8.5, tl=tl, at=1550)
        dv = self.deco(s, 'divider', 1.1, 4.95, w=4.6)
        it = add_text(s, 1.14, 5.3, 5.6, 0.7, items.split('\n'), size=11.5, color='ink_soft', font=F['body'], line=1.5) if items else None
        pm = add_text(s, 1.14, 6.45, 5.6, 0.35, poem, size=12, color='line', latin=F['latin'], italic=True) if poem else None
        fl = self.floaters()
        spots = [(6.15, 0.95, 0.42), (6.75, 6.1, 0.38), (5.3, 2.1, 0.3)]
        leaves = [self.deco(s, os.path.join(self.theme_dir, fl[k]), *spots[k][:2], w=spots[k][2]) for k in range(min(3, len(fl)))]
        tl.add(p, 'zoom', at=0, dur=1300, scale=0.92)
        for i, f in enumerate(fr):
            tl.add(f, 'fade', at=400 + i * 200, dur=900)
        tl.add(wm, 'fade', at=900, dur=1200); tl.add(ob, 'rise', at=800, dur=1100, dy=-0.05); tl.add(q, 'fade', at=0, dur=1200)
        tl.add(big, 'fade', at=200, dur=1200); tl.add(pt, 'rise', at=300, dur=900)
        tl.add(t, 'letters', at=600, dur=500, pct=18000); tl.add(et, 'wipe', at=1300, dur=700)
        tl.add(dv, 'split', at=1750, dur=600); tl.add(it, 'fade', at=1950, dur=700); tl.add(pm, 'fade', at=2250, dur=900)
        for k, l in enumerate(leaves):
            tl.add(l, 'rise', at=1200 + k * 250, dur=1200, dy=-0.08)
        return self.finish(s, tl, note or f'第{num}部分：{cn}。')

    def overview(self, title, en, stats, facts, paragraph, photos, note=''):
        """stats: [(big, unit, caption)]x3 ; facts: [(k, v)] ; photos: [(path, focus, label)]x3"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en)
        for i, (big, unit, cap) in enumerate(stats[:3]):
            x = 1.05 + i * 2.0
            c = self.card(s, x, 1.62, 1.82, 1.25)
            bt = add_text(s, x, 1.74, 1.82, 0.62, [[dict(text=big, size=19 if len(big) > 6 else (22 if len(big) > 2 else 26), bold=True),
                                                     dict(text=unit, size=11, color='ink_soft')]], color='accent', font=F['head'], latin=F['latin'], align='c', anchor='m')
            ct = add_text(s, x, 2.4, 1.82, 0.3, cap, size=10, color='ink_soft', font=F['body'], align='c')
            at = 800 + i * 200
            tl.add(c, 'zoom', at=at, dur=500, scale=0.85); tl.add(bt, 'zoom', at=at + 100, dur=500, scale=0.6); tl.add(ct, 'fade', at=at + 300, dur=500)
        self.bullets(s, 1.07, 3.2, 5.9, facts, tl, 1500)
        para = add_text(s, 1.07, 3.2 + 0.42 * len(facts) + 0.1, 5.85, 1.9, paragraph, size=12, color='ink_soft', font=F['body'], line=1.45)
        tl.add(para, 'fade', at=2200, dur=900)
        xs, ys = [7.35, 9.2, 11.05], [1.85, 2.45, 1.65]
        for i, (ph, foc, lab) in enumerate(photos[:3]):
            pic, fr = self.framed_photo(s, ph, xs[i], ys[i], 1.62, 3.95, frame=0.1, focus=foc)
            lt = add_text(s, xs[i], ys[i] + 4.12, 1.62, 0.3, lab, size=11, color='primary_dk', font=F['head'], bold=True, align='c')
            at = 300 + i * 250
            tl.add(pic, 'rise', at=at, dur=1000, dy=0.08)
            for f in fr:
                tl.add(f, 'fade', at=at + 400, dur=800)
            tl.add(lt, 'fade', at=at + 700, dur=500)
        return self.finish(s, tl, note)

    def text_image(self, title, en, heading, paragraphs, photo, focus=(0.5, 0.5), items=None, note=''):
        """generic: heading + paragraphs (+ optional (k,v) bullet rows) left, framed photo right"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        h = add_text(s, 1.05, 1.6, 6.2, 0.6, heading, size=22, color='primary_dk', font=F['head'], bold=True)
        dv = self.deco(s, 'divider', 1.0, 2.3, w=4.2)
        y = 2.65
        if items:
            self.bullets(s, 1.07, y, 6.0, items, tl, 900)
            y += 0.42 * len(items) + 0.15
        para = add_text(s, 1.07, y, 6.0, 6.9 - y, paragraphs, size=12.5, color='ink_soft', font=F['body'], line=1.5, para_after=6)
        p, fr = self.framed_photo(s, photo, 8.0, 1.55, 4.3, 5.3, frame=0.13, focus=focus)
        tl.add(h, 'wipe', at=400, dur=700); tl.add(dv, 'split', at=700, dur=600); tl.add(para, 'fade', at=1300, dur=900)
        tl.add(p, 'rise', at=300, dur=1000, dy=0.06)
        for f in fr:
            tl.add(f, 'fade', at=700, dur=800)
        return self.finish(s, tl, note)

    def gallery_grid(self, title, en, tiles, cols=4, note=''):
        """tiles: [(photo, caption)] up to 2 rows; framed with theme geometry"""
        s, tl, n = self.new('gallery')
        self.chrome(s, n, tl, title, en, sprig=False)
        rows = (len(tiles) + cols - 1) // cols
        gw = 11.4 / cols
        th = (5.3 - 0.45 * rows) / rows
        for i, (ph, cap) in enumerate(tiles):
            r, c = divmod(i, cols)
            x, y = 0.98 + c * gw, 1.55 + r * (th + 0.45)
            pic = add_pic(s, self.A(ph), x, y, gw - 0.25, th, geom='frame')
            shadow(pic, blur=0.2, dist=0.05, alpha=25)
            t = add_text(s, x, y + th + 0.06, gw - 0.25, 0.3, cap, size=11, color='primary_dk', font=F['head'], bold=True, align='c')
            tl.add(pic, 'rise', at=400 + i * 140, dur=800, dy=0.05); tl.add(t, 'fade', at=700 + i * 140, dur=500)
        return self.finish(s, tl, note)

    def location_levels(self, title, en, levels, callout=None, note=''):
        """levels: [(img, title, text)]x3 (e.g. country›province›city maps from scripts/three)
        callout: dict(meta='3d/map_chengdu_meta.json', pin='武侯区', text='项目所在 · 武侯区') on the 3rd map"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        boxes = [(0.75, 1.75, 3.7, 2.75), (4.75, 1.55, 3.55, 3.0), (8.55, 1.35, 4.3, 3.35)]
        for i, (rel, ttl, txt) in enumerate(levels[:3]):
            at = 500 + i * 650
            if self.has(rel):
                mx, my, pw, ph, sc = self.fit(rel, boxes[i])
                pic = add_pic(s, self.A(rel), mx, my, pw, ph, cover=False)
                tl.add(pic, 'rise', at=at, dur=1000, dy=0.06)
                if callout and i == 2:
                    hx, hy = json.load(open(self.A(callout['meta']), encoding='utf8'))['pins'][callout['pin']]
                    ax, ay = mx + hx * sc, my + hy * sc
                    ld = add_line(s, ax, ay, ax + 0.75, ay - 0.62, color='accent', width=1)
                    dot = add_shape(s, 'oval', ax - 0.05, ay - 0.05, 0.1, 0.1, fill='accent')
                    pl = add_shape(s, 'pill', ax + 0.72, ay - 0.82, 1.55, 0.36, fill='accent')
                    ptx = add_text(s, ax + 0.72, ay - 0.82, 1.55, 0.36, callout['text'], size=9.5, color='white', font=F['head'], bold=True, align='c', anchor='m')
                    tl.add(dot, 'zoom', at=at + 900, dur=300); tl.add(ld, 'wipe', at=at + 1100, dur=400)
                    tl.add(pl, 'zoom', at=at + 1400, dur=400); tl.add(ptx, 'zoom', at=at + 1400, dur=400)
            cx = [2.6, 6.52, 10.7][i]
            b = self.deco(s, 'bullet', cx - 0.13, 4.95, w=0.26)
            tt = add_text(s, cx - 1.9, 5.25, 3.8, 0.42, ttl, size=16, color='primary_dk', font=F['head'], bold=True, align='c')
            dt = add_text(s, cx - 1.65, 5.72, 3.3, 1.2, txt, size=10.5, color='ink_soft', font=F['body'], line=1.35, align='c')
            tl.add(b, 'zoom', at=at + 400, dur=400); tl.add(tt, 'fade', at=at + 450, dur=600); tl.add(dt, 'fade', at=at + 650, dur=700)
        for i, (x1, x2) in enumerate([(4.35, 5.05), (8.25, 8.95)]):
            a1 = add_line(s, x1, 3.05, x2, 3.05, color='line', width=1.5, dash='sysDot')
            tl.add(a1, 'wipe', at=1000 + i * 650, dur=500)
        return self.finish(s, tl, note)

    def poi_map(self, title, en, map_img, meta, pois, center=None, neighbors=(), insight=None, nudge=None, note=''):
        """pois: [(icon, name, desc)] — name must match meta['pins'] keys; numbered badges sit on pin heads.
        center: (big_label, en_label) at meta['labels'][first key]; neighbors: [(text, fx, fy)] fractions of map box.
        insight: (title, en, text) olive card under the list."""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        m = json.load(open(self.A(meta), encoding='utf8'))
        px0, py0, pw, ph, sc = self.fit(map_img, (0.75, 1.75, 6.9, 5.15))
        pic = add_pic(s, self.A(map_img), px0, py0, pw, ph, cover=False)
        tl.add(pic, 'zoom', at=300, dur=1100, scale=0.9)
        nudge = nudge or {}
        for k, (_, name, _) in enumerate(pois):
            if name not in m.get('pins', {}):
                continue
            hx, hy = m['pins'][name]; dx, dy = nudge.get(name, (0, 0))
            x, y = px0 + hx * sc + dx, py0 + hy * sc + dy
            c = add_shape(s, 'oval', x - 0.13, y - 0.13, 0.26, 0.26, fill='accent', line='white', lw=1)
            ct = add_text(s, x - 0.13, y - 0.13, 0.26, 0.26, str(k + 1), size=9, color='white', latin=F['latin'], bold=True, align='c', anchor='m')
            tl.add(c, 'zoom', at=1400 + k * 160, dur=400); tl.add(ct, 'zoom', at=1400 + k * 160, dur=400)
        if center and m.get('labels'):
            cx, cy = list(m['labels'].values())[0]
            wl = add_text(s, px0 + cx * sc - 1.3, py0 + cy * sc - 0.35, 2.6, 0.7,
                          [[dict(text=center[0], size=24, bold=True, color='deep')], [dict(text=center[1], latin=True, italic=True, size=8.5, color='line', spacing=2)]],
                          font=F['head'], align='c')
            tl.add(wl, 'fade', at=1200, dur=800)
        for k, (txt, fx, fy) in enumerate(neighbors):
            t = add_text(s, px0 + pw * fx, py0 + ph * fy, 1.9, 0.28, txt, size=10.5, color='primary_dk', font=F['head'], bold=True)
            tl.add(t, 'fade', at=1300 + k * 150, dur=600)
        tl.add(self.icon(s, 'compass', 0.85, 6.35, 0.42, 'primary'), 'zoom', at=1300, dur=500)
        for k, (ic, name, desc) in enumerate(pois):
            y = 1.55 + k * 0.66
            for b in self.icon_badge(s, ic, 8.05, y, d=0.46, ring=False):
                tl.add(b, 'zoom', at=900 + k * 200, dur=450)
            nt = add_text(s, 8.7, y - 0.04, 4.0, 0.3, [[dict(text=f'{k + 1:02d}  ', color='accent', latin=True, bold=True), dict(text=name)]],
                          size=13, color='ink', font=F['head'], bold=True)
            dt = add_text(s, 8.7, y + 0.27, 4.0, 0.28, desc, size=10.5, color='ink_soft', font=F['body'])
            tl.add(nt, 'slide', at=980 + k * 200, dur=600, dx=0.03); tl.add(dt, 'fade', at=1150 + k * 200, dur=600)
        if insight:
            yb = 1.55 + len(pois) * 0.66 + 0.1
            c = self.card(s, 8.0, yb, 4.75, 6.9 - yb, fill='primary', alpha=92, line=None)
            ct = add_text(s, 8.25, yb + 0.15, 4.3, 6.9 - yb - 0.3,
                          [[dict(text=insight[0] + '　', bold=True, font=F['head'], size=13), dict(text=insight[1], latin=True, italic=True, size=8.5, color='line_lt')],
                           insight[2]], size=11, color='white', font=F['body'], line=1.4, para_after=4)
            tl.add(c, 'wipe', at=2300, dur=700); tl.add(ct, 'fade', at=2700, dur=800)
        return self.finish(s, tl, note)

    def chart_page(self, title, en, line_chart, bar_chart, stats, responses, source='', resp_title=('设计回应', 'DESIGN RESPONSE'), note=''):
        """line_chart / bar_chart: (title, categories, values, number_format) ; stats: [(icon, value, unit, caption)]x4 ;
        responses: [(key, text)]"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        cc = self.card(s, 0.95, 1.55, 6.55, 5.35, fill='white', alpha=45)
        tl.add(cc, 'fade', at=300, dur=700)

        def style(ch, ttl):
            ch.has_title = True; ch.chart_title.text_frame.text = ttl
            r = ch.chart_title.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = rgb('primary_dk'); set_fonts(r._r.get_or_add_rPr(), F['head'], F['head'])
            ch.has_legend = False; ch.font.size = Pt(9); ch.font.color.rgb = rgb('ink_soft'); ch.font.name = F['body']
            va = ch.value_axis; va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb('line_lt'); va.major_gridlines.format.line.width = Pt(0.5)
            va.format.line.fill.background(); va.tick_labels.font.size = Pt(8.5)
            ca = ch.category_axis; ca.format.line.color.rgb = rgb('line_lt'); ca.tick_labels.font.size = Pt(8.5)
        if line_chart:
            ttl, cats, vals, fmt = (list(line_chart) + ['0.0'])[:4]
            cd = CategoryChartData(); cd.categories = cats; cd.add_series(ttl, vals)
            gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, E(1.1), E(1.65), E(6.25), E(2.55), cd)
            ch = gf.chart; style(ch, ttl)
            ser = ch.plots[0].series[0]
            ser.format.line.color.rgb = rgb('accent'); ser.format.line.width = Pt(2.25); ser.smooth = True
            ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 6
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb('accent'); ser.marker.format.line.color.rgb = rgb('white')
            pl = ch.plots[0]; pl.has_data_labels = True; dl = pl.data_labels
            dl.font.size = Pt(8); dl.font.color.rgb = rgb('accent'); dl.position = XL_LABEL_POSITION.ABOVE
            dl.number_format = fmt; dl.number_format_is_linked = False
            tl.add(gf, 'wipe', at=700, dur=1000)
        if bar_chart:
            ttl, cats, vals = bar_chart[:3]
            cd = CategoryChartData(); cd.categories = cats; cd.add_series(ttl, vals)
            gf2 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, E(1.1), E(4.25), E(6.25), E(2.55), cd)
            ch2 = gf2.chart; style(ch2, ttl)
            ser2 = ch2.plots[0].series[0]; ser2.format.fill.solid(); ser2.format.fill.fore_color.rgb = rgb('primary_lt')
            ch2.plots[0].gap_width = 60; ch2.plots[0].has_data_labels = True
            d2 = ch2.plots[0].data_labels; d2.font.size = Pt(8); d2.font.color.rgb = rgb('primary_dk'); d2.position = XL_LABEL_POSITION.OUTSIDE_END
            tl.add(gf2, 'wipe', at=1100, dur=1000)
        if source:
            tl.add(add_text(s, 1.0, 6.93, 6.5, 0.25, source, size=8.5, color='ink_soft', font=F['body']), 'fade', at=1600, dur=600)
        for i, (ic, v, u, cap) in enumerate(stats[:4]):
            x = 7.9 + (i % 2) * 2.45; y = 1.6 + (i // 2) * 1.12
            c = self.card(s, x, y, 2.25, 0.95, lw=0.5)
            ico = self.icon(s, ic, x + 0.15, y + 0.25, 0.44, 'accent')
            vt = add_text(s, x + 0.7, y + 0.1, 1.5, 0.5, [[dict(text=v, size=20, bold=True), dict(text=u, size=10, color='ink_soft')]],
                          color='accent', latin=F['latin'], font=F['head'])
            ctt = add_text(s, x + 0.72, y + 0.6, 1.5, 0.26, cap, size=9.5, color='ink_soft', font=F['body'])
            at = 900 + i * 150
            tl.add(c, 'zoom', at=at, dur=450, scale=0.85); tl.add(ico, 'zoom', at=at + 100, dur=400)
            tl.add(vt, 'fade', at=at + 150, dur=500); tl.add(ctt, 'fade', at=at + 250, dur=500)
        self.side_heading(s, 7.92, 3.95, resp_title[0], resp_title[1], tl, 1700)
        y = 4.42
        for i, (k, v) in enumerate(responses):
            b = self.deco(s, 'node', 7.95, y + 0.05, w=0.2)
            t = add_text(s, 8.25, y, 4.5, 0.58, [[dict(text=k + '　', color='accent', font=F['head'], bold=True), dict(text=v)]],
                         size=10.5, color='ink_soft', font=F['body'], line=1.25)
            tl.add(b, 'zoom', at=1900 + i * 180, dur=400); tl.add(t, 'slide', at=1900 + i * 180, dur=600, dx=0.03)
            y += 0.58 if len(k) + len(v) > 30 else 0.4     # 2-line vs 1-line rows keep even gaps
        return self.finish(s, tl, note)

    def persona(self, title, en, people, extra_line=None, day=(), mapping=(), mapping_title='需求 → 空间回应', note=''):
        """people: [(icon, name, role, hobby, need)]x3 ; extra_line: (key, text) ; day: [(time, act, room)]<=6 ;
        mapping: [(need, space)]<=6"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        fills = ['primary', 'accent', 'extra']
        for i, (ic, name, role, hobby, need) in enumerate(people[:3]):
            x = 0.95 + i * 2.62
            c = self.card(s, x, 1.95, 2.4, 3.05, kind='frame')
            bd = self.icon_badge(s, ic, x + 0.86, 1.62, d=0.68, fill=fills[i])
            nt = add_text(s, x, 2.48, 2.4, 0.35, name, size=14, color='ink', font=F['head'], bold=True, align='c')
            rt = add_text(s, x, 2.85, 2.4, 0.28, role, size=10, color='line', font=F['body'], align='c')
            ht = add_text(s, x + 0.2, 3.3, 2.0, 1.6, [[dict(text='爱好', color='primary_dk', font=F['head'], bold=True)], hobby,
                                                      [dict(text='需求', color='primary_dk', font=F['head'], bold=True)], need],
                          size=10.5, color='ink_soft', font=F['body'], align='c', line=1.25, para_after=2)
            at = 600 + i * 300
            tl.add(c, 'rise', at=at, dur=900, dy=0.06)
            for b in bd:
                tl.add(b, 'zoom', at=at + 300, dur=450)
            tl.add(nt, 'fade', at=at + 400, dur=500); tl.add(rt, 'fade', at=at + 500, dur=500); tl.add(ht, 'fade', at=at + 600, dur=600)
        if extra_line:
            tl.add(add_text(s, 0.95, 5.25, 7.6, 0.3, [[dict(text=extra_line[0] + '　', color='primary_dk', bold=True, font=F['head']),
                                                       dict(text=extra_line[1])]], size=11, color='ink_soft', font=F['body']), 'fade', at=1700, dur=600)
        if day:
            self.side_heading(s, 0.95, 5.72, '一日生活', 'A DAY AT HOME', tl, 1900)
            ln = add_line(s, 1.05, 6.38, 8.45, 6.38, color='line', width=0.9)
            tl.add(ln, 'wipe', at=2100, dur=900)
            step = 7.4 / max(len(day) - 1, 1)
            for i, (tm, act, rm) in enumerate(day):
                x = 1.05 + i * step
                d = add_shape(s, 'oval', x - 0.07, 6.31, 0.14, 0.14, fill='accent')
                tt = add_text(s, x - 0.6, 6.0, 1.2, 0.26, tm, size=10, color='accent', latin=F['latin'], bold=True, align='c')
                a_ = add_text(s, x - 0.7, 6.5, 1.4, 0.45, [act, [dict(text=rm, size=9, color='line')]], size=10, color='ink', font=F['body'], align='c', line=1.1)
                at = 2200 + i * 140
                tl.add(d, 'zoom', at=at, dur=350); tl.add(tt, 'fade', at=at, dur=400); tl.add(a_, 'fade', at=at + 100, dur=450)
        if mapping:
            rc = self.card(s, 8.85, 1.55, 3.9, 5.35, fill='primary', alpha=93, line=None)
            tl.add(rc, 'wipe', at=400, dur=800)
            tl.add(add_text(s, 9.1, 1.72, 3.5, 0.4, mapping_title, size=14, color='white', font=F['head'], bold=True), 'fade', at=900, dur=500)
            for i, (a, b) in enumerate(mapping[:6]):
                y = 2.3 + i * 0.75
                t1 = add_text(s, 9.1, y, 3.5, 0.3, a, size=11, color='line_lt', font=F['head'], bold=True)
                t2 = add_text(s, 9.1, y + 0.3, 3.5, 0.3, b, size=11, color='white', font=F['body'])
                at = 1100 + i * 170
                tl.add(t1, 'slide', at=at, dur=500, dx=0.03); tl.add(t2, 'fade', at=at + 100, dur=500)
                if i < len(mapping) - 1:
                    tl.add(add_line(s, 9.1, y + 0.66, 12.45, y + 0.66, color='line_lt', width=0.4, alpha=50), 'wipe', at=at + 150, dur=400)
        return self.finish(s, tl, note)

    def pillars(self, title, en, quote, cols, note=''):
        """quote: str ; cols: [(char, title, en, photo, text)]x3 — concept / strategy pillars"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en)
        q = add_text(s, 1.0, 1.5, 11.3, 0.5, [[dict(text='“ ', latin=True, color='line', size=22), dict(text=quote),
                                               dict(text=' ”', latin=True, color='line', size=22)]], size=17, color='ink', font=F['head'], align='c', anchor='m')
        tl.add(q, 'fade', at=500, dur=900)
        fills = ['primary', 'accent', 'accent2']
        for i, (ch_, ttl, ens, ph, txt) in enumerate(cols[:3]):
            x = 1.0 + i * 3.85
            c = self.card(s, x, 2.3, 3.45, 4.6, alpha=60)
            pic = add_pic(s, self.A(ph), x + 0.2, 2.5, 3.05, 1.85, focus=(0.5, 0.4))
            bd = add_shape(s, 'oval', x + 1.27, 3.93, 0.9, 0.9, fill=fills[i], line='paper', lw=2)
            bt = add_text(s, x + 1.27, 3.93, 0.9, 0.9, ch_, size=26, color='white', font=F['head'], bold=True, align='c', anchor='m')
            tt = add_text(s, x, 4.92, 3.45, 0.4, [[dict(text=ttl + '  ', bold=True), dict(text=ens, latin=True, italic=True, size=9, color='line')]],
                          size=15, color='primary_dk', font=F['head'], align='c')
            dt = add_text(s, x + 0.3, 5.38, 2.85, 1.4, txt, size=11, color='ink_soft', font=F['body'], line=1.4, align='c')
            at = 900 + i * 350
            tl.add(c, 'rise', at=at, dur=900, dy=0.05); tl.add(pic, 'fade', at=at + 200, dur=800)
            tl.add(bd, 'zoom', at=at + 500, dur=450); tl.add(bt, 'zoom', at=at + 500, dur=450)
            tl.add(tt, 'fade', at=at + 650, dur=500); tl.add(dt, 'fade', at=at + 800, dur=600)
        return self.finish(s, tl, note)

    def moodboard(self, title, en, heading_lines, heading_en, text, keywords, tiles, note=''):
        """heading_lines: [(text, color)] ; keywords: [str]<=6 ; tiles: [(photo, title, caption)]x8 (4x2 framed)"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        hd = add_text(s, 1.0, 1.6, 3.6, 1.0, [[dict(text=t, color=c)] for t, c in heading_lines], size=26, color='primary_dk', font=F['head'], bold=True, line=1.1)
        sub = add_text(s, 1.02, 2.72, 3.6, 0.3, heading_en, size=9, color='line', latin=F['latin'], italic=True, spacing=2)
        dv = self.deco(s, 'divider', 1.0, 3.12, w=3.3)
        txt = add_text(s, 1.02, 3.42, 3.45, 2.2, text, size=11.5, color='ink_soft', font=F['body'], line=1.45)
        tl.add(hd, 'wipe', at=400, dur=700); tl.add(sub, 'wipe', at=700, dur=600); tl.add(dv, 'split', at=900, dur=600); tl.add(txt, 'fade', at=1100, dur=800)
        for i, kw in enumerate(keywords[:6]):
            x = 1.02 + (i % 3) * 1.15; y = 5.75 + (i // 3) * 0.5
            c = add_shape(s, 'pill', x, y, 1.0, 0.36, fill='primary' if i % 2 == 0 else 'accent', alpha=90)
            t = add_text(s, x, y, 1.0, 0.36, kw, size=11, color='white', font=F['head'], bold=True, align='c', anchor='m')
            tl.add(c, 'zoom', at=1800 + i * 120, dur=400); tl.add(t, 'zoom', at=1800 + i * 120, dur=400)
        for i, (ph, ttl, cap) in enumerate(tiles[:8]):
            r, c_ = divmod(i, 4)
            x = 4.95 + c_ * 2.0; y = 1.5 + r * 2.82
            pic = add_pic(s, self.A(ph), x, y, 1.72, 1.95, geom='frame')
            shadow(pic, blur=0.2, dist=0.05, alpha=25)
            fr = add_shape(s, 'frame', x - 0.07, y - 0.07, 1.86, 2.09, line='line', lw=0.6)
            tt = add_text(s, x - 0.1, y + 2.04, 1.92, 0.3, ttl, size=12, color='primary_dk', font=F['head'], bold=True, align='c')
            ct = add_text(s, x - 0.1, y + 2.31, 1.92, 0.25, cap, size=9, color='ink_soft', font=F['body'], align='c')
            at = 500 + i * 160
            tl.add(pic, 'rise', at=at, dur=800, dy=0.05); tl.add(fr, 'fade', at=at + 300, dur=600)
            tl.add(tt, 'fade', at=at + 400, dur=500); tl.add(ct, 'fade', at=at + 500, dur=500)
        return self.finish(s, tl, note)

    def palette(self, title, en, statement, swatches, ratio, ratio_note, photo, focus=(0.5, 0.5), note=''):
        """swatches: [(hex, name, use)]<=8 ; ratio: [(fraction, hex, label, text_color)] ; photo framed on the right"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        pos = add_text(s, 1.0, 1.55, 7.9, 0.6, [[dict(text='定位　', color='accent', bold=True, font=F['head']), dict(text=statement)]],
                       size=11.5, color='ink_soft', font=F['body'], line=1.4)
        tl.add(pos, 'fade', at=400, dur=800)
        step = 7.9 / max(len(swatches), 1)
        for i, (hx, nm, use) in enumerate(swatches[:8]):
            x = 1.0 + i * step
            light = sum(int(hx[j:j + 2], 16) for j in (0, 2, 4)) > 600
            sw = add_shape(s, 'frame', x, 2.45, step - 0.15, 1.75, fill=hx, line='line' if light else None, lw=0.5)
            shadow(sw, blur=0.12, dist=0.04, alpha=22)
            nt = add_text(s, x - 0.1, 4.3, step + 0.05, 0.3, nm, size=10.5, color='ink', font=F['head'], bold=True, align='c')
            ht = add_text(s, x - 0.1, 4.58, step + 0.05, 0.22, '#' + hx, size=8, color='ink_soft', latin=F['latin'], align='c')
            ut = add_text(s, x - 0.12, 4.8, step + 0.09, 0.4, use, size=8.5, color='line', font=F['body'], align='c')
            at = 700 + i * 130
            tl.add(sw, 'wipe', at=at, dur=600, dir='up'); tl.add(nt, 'fade', at=at + 250, dur=450)
            tl.add(ht, 'fade', at=at + 300, dur=450); tl.add(ut, 'fade', at=at + 350, dur=450)
        hd = add_text(s, 1.0, 5.42, 7, 0.3, [[dict(text='配色比例　', bold=True), dict(text=' · '.join(str(int(f * 100)) for f, *_ in ratio), latin=True, italic=True, color='line', bold=False)]],
                      size=13, color='primary_dk', font=F['head'])
        tl.add(hd, 'wipe', at=1900, dur=500)
        x = 1.0
        for i, (f, hx, lb, tc) in enumerate(ratio):
            w = 7.85 * f
            r = add_shape(s, 'rect', x, 5.85, w - 0.03, 0.42, fill=hx)
            t = add_text(s, x + 0.12, 5.85, w - 0.2, 0.42, lb, size=9.5 if f > 0.15 else 8.5, color=tc, font=F['body'], anchor='m')
            tl.add(r, 'wipe', at=2100 + i * 250, dur=500); tl.add(t, 'fade', at=2250 + i * 250, dur=400)
            x += w
        tl.add(add_text(s, 1.0, 6.4, 7.9, 0.4, ratio_note, size=9.5, color='ink_soft', font=F['body']), 'fade', at=2900, dur=600)
        p, fr = self.framed_photo(s, photo, 9.55, 1.6, 3.0, 5.1, frame=0.12, focus=focus)
        tl.add(p, 'rise', at=300, dur=1000, dy=0.06)
        for f_ in fr:
            tl.add(f_, 'fade', at=700, dur=800)
        return self.finish(s, tl, note)

    def materials(self, title, en, mats, subs=(), note=''):
        """mats: [(texture_img, application_img|None, name, rooms, text)]x3 ; subs: [(img, name, use)]<=4"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        for i, (m, u, nm, rm, txt) in enumerate(mats[:3]):
            cx = 2.45 + i * 3.6
            ring = add_shape(s, 'oval', cx - 1.08, 1.47, 2.16, 2.16, line='line', lw=0.8)
            mp = add_pic(s, self.A(m), cx - 0.98, 1.57, 1.96, 1.96, geom='oval')
            shadow(mp, blur=0.2, dist=0.05, alpha=25)
            up = None
            if u and self.has(u):
                up = add_pic(s, self.A(u), cx + 0.45, 2.75, 1.05, 1.05, geom='oval')
                pic_outline(up, 'paper', 2.5); shadow(up, blur=0.15, dist=0.04, alpha=30)
            nt = add_text(s, cx - 1.5, 3.88, 3.0, 0.35, nm, size=14, color='primary_dk', font=F['head'], bold=True, align='c')
            rt = add_text(s, cx - 1.5, 4.22, 3.0, 0.26, rm, size=9.5, color='line', font=F['body'], align='c')
            dt = add_text(s, cx - 1.45, 4.5, 2.9, 0.75, txt, size=10.5, color='ink_soft', font=F['body'], align='c', line=1.35)
            at = 500 + i * 300
            tl.add(ring, 'zoom', at=at, dur=600); tl.add(mp, 'circle', at=at + 150, dur=700); tl.add(up, 'zoom', at=at + 600, dur=450)
            tl.add(nt, 'fade', at=at + 500, dur=500); tl.add(rt, 'fade', at=at + 600, dur=500); tl.add(dt, 'fade', at=at + 700, dur=600)
        if subs:
            tl.add(self.deco(s, 'divider', 3.0, 5.35, w=7.3), 'split', at=1700, dur=700)
            for i, (ph, nm, us) in enumerate(subs[:4]):
                x = 1.05 + i * 2.95
                c = add_pic(s, self.A(ph), x, 5.82, 0.82, 0.82, geom='oval')
                pic_outline(c, 'line', 0.75)
                t = add_text(s, x + 0.98, 5.9, 1.9, 0.68, [[dict(text=nm, font=F['head'], bold=True, color='ink', size=12)], [dict(text=us, size=9.5, color='ink_soft')]],
                             font=F['body'], line=1.2)
                tl.add(c, 'circle', at=2000 + i * 180, dur=550); tl.add(t, 'slide', at=2150 + i * 180, dur=600, dx=0.03)
        return self.finish(s, tl, note)

    def _plan_frame(self, s, x, y, w, h, tl, at=200):
        fr = add_shape(s, 'rect', x - 0.1, y - 0.1, w + 0.2, h + 0.2, fill='white', alpha=55, line='line', lw=0.6)
        shadow(fr, blur=0.25, dist=0.05, alpha=16)
        tl.add(fr, 'fade', at=at, dur=700)

    def plan_annotated(self, title, en, plan_img, rooms, dims=(), points=(), offset=(0, 0), note=''):
        """plan_img: line plan ; rooms: [(name, (px,py))] numbered badges at source-pixel positions
        (minus offset if the image was cropped) ; dims: [(value, caption)]<=4 ; points: ['key：text']<=3 (one line each)"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        x, y, w, h, sc = self.fit(plan_img, (0.95, 1.45, 6.4, 5.65))
        size = Image.open(self.A(plan_img)).size
        self._plan_frame(s, x, y, w, h, tl)
        tl.add(add_pic(s, self.A(plan_img), x, y, w, h, cover=False), 'wipe', at=400, dur=1200, dir='down')
        for i, (nm, (px, py)) in enumerate(rooms):
            bx, by = self.pmap(px, py, (x, y, w, h), size, offset)
            c = add_shape(s, 'oval', bx - 0.15, by - 0.15, 0.3, 0.3, fill='accent', line='white', lw=0.75)
            t = add_text(s, bx - 0.15, by - 0.15, 0.3, 0.3, f'{i + 1}', size=9, color='white', latin=F['latin'], bold=True, align='c', anchor='m')
            tl.add(c, 'zoom', at=1500 + i * 90, dur=350); tl.add(t, 'zoom', at=1500 + i * 90, dur=350)
        self.side_heading(s, 7.85, 1.5, '空间构成', 'Spaces', tl, 600)
        for i, (nm, _) in enumerate(rooms):
            cx = 7.88 + (i % 3) * 1.62; cy = 1.98 + (i // 3) * 0.42
            tl.add(add_text(s, cx, cy, 1.6, 0.32, [[dict(text=f'{i + 1:02d} ', color='accent', latin=True, bold=True), dict(text=nm)]], size=11.5, color='ink', font=F['body']),
                   'fade', at=800 + i * 70, dur=400)
        y2 = 1.98 + ((len(rooms) + 2) // 3) * 0.42 + 0.1
        if dims:
            self.side_heading(s, 7.85, y2, '主要尺寸', 'Dimensions', tl, 1500)
            for i, (v, cap) in enumerate(dims[:4]):
                dx = 7.88 + i * 1.22
                c = self.card(s, dx, y2 + 0.42, 1.1, 0.82, lw=0.5)
                vt = add_text(s, dx, y2 + 0.49, 1.1, 0.4, v, size=15, color='accent', latin=F['latin'], bold=True, align='c')
                ct = add_text(s, dx, y2 + 0.88, 1.1, 0.28, cap, size=8.5, color='ink_soft', font=F['body'], align='c')
                tl.add(c, 'zoom', at=1700 + i * 120, dur=400, scale=0.85); tl.add(vt, 'fade', at=1800 + i * 120, dur=400); tl.add(ct, 'fade', at=1850 + i * 120, dur=400)
            y2 += 1.47
        if points:
            self.side_heading(s, 7.85, y2, '布局策略', 'Planning', tl, 2200)
            for i, ptx in enumerate(points[:3]):
                yy = y2 + 0.43 + i * 0.4
                b = self.deco(s, 'node', 7.9, yy + 0.06, w=0.17)
                k, v = ptx.split('：', 1) if '：' in ptx else ('', ptx)
                t = add_text(s, 8.15, yy, 4.6, 0.4, [[dict(text=k + '  ', color='primary_dk', font=F['head'], bold=True), dict(text=v)]], size=9.5, color='ink_soft', font=F['body'], line=1.15)
                tl.add(b, 'zoom', at=2400 + i * 160, dur=350); tl.add(t, 'slide', at=2400 + i * 160, dur=550, dx=0.03)
        return self.finish(s, tl, note)

    def plan_legend(self, title, en, plan_img, legend, note_text='', note=''):
        """legend: [(swatch_img, name, rooms)]<=5 ; note_text: card at bottom right"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        x, y, w, h, sc = self.fit(plan_img, (0.95, 1.45, 6.4, 5.65))
        self._plan_frame(s, x, y, w, h, tl)
        tl.add(add_pic(s, self.A(plan_img), x, y, w, h, cover=False), 'zoom', at=400, dur=1100, scale=0.94)
        self.side_heading(s, 7.85, 1.5, '图例', 'Legend', tl, 700)
        for i, (ph, nm, rm) in enumerate(legend[:5]):
            yy = 2.0 + i * 0.78
            c = add_pic(s, self.A(ph), 7.9, yy, 0.6, 0.6, geom='oval') if self.has(ph) else add_shape(s, 'oval', 7.9, yy, 0.6, 0.6, fill=ph)
            if self.has(ph):
                pic_outline(c, 'line', 0.75)
            t = add_text(s, 8.7, yy + 0.02, 4.0, 0.6, [[dict(text=nm, font=F['head'], bold=True, color='ink', size=12.5)], [dict(text=rm, size=9.5, color='line')]], font=F['body'], line=1.15)
            tl.add(c, 'circle', at=900 + i * 160, dur=500); tl.add(t, 'slide', at=1000 + i * 160, dur=550, dx=0.03)
        if note_text:
            c = self.card(s, 7.85, 5.95, 4.9, 0.95, fill='primary', alpha=92, line=None)
            t = add_text(s, 8.05, 6.03, 4.55, 0.8, note_text, size=10, color='white', font=F['body'], line=1.35, anchor='m')
            tl.add(c, 'wipe', at=1900, dur=600); tl.add(t, 'fade', at=2200, dur=700)
        return self.finish(s, tl, note)

    def _labels_on(self, s, tl, meta_labels, names, box_xy, sc, colors=None, at=1500):
        x, y = box_xy
        for i, nm in enumerate(names):
            if nm not in meta_labels:
                continue
            lx, ly = meta_labels[nm]
            cx, cy = x + lx * sc, y + ly * sc
            tw = 0.22 + 0.17 * len(nm.replace('/', ''))
            fill, tc = (colors[nm], 'white') if colors else ('paper', 'ink')
            pill = add_shape(s, 'pill', cx - tw / 2, cy - 0.16, tw, 0.3, fill=fill, alpha=92, line='white' if colors else 'line', lw=0.5)
            if not colors:
                shadow(pill, blur=0.08, dist=0.02, alpha=25)
            t = add_text(s, cx - tw / 2, cy - 0.16, tw, 0.3, nm, size=9.5, color=tc, font=F['head'], bold=True, align='c', anchor='m')
            tl.add(pill, 'zoom', at=at + i * 90, dur=400); tl.add(t, 'zoom', at=at + i * 90, dur=400)

    def model3d(self, title, en, img, meta, labels, intro, points, note=''):
        """3D render (scripts/three) with projected room labels ; points: [(icon, key, text)]x3"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        m = json.load(open(self.A(meta), encoding='utf8'))
        x, y, w, h, sc = self.fit(img, (0.75, 1.35, 7.9, 5.95))
        tl.add(add_pic(s, self.A(img), x, y, w, h, cover=False), 'rise', at=300, dur=1300, dy=0.06)
        self._labels_on(s, tl, m['labels'], labels, (x, y), sc)
        self.side_heading(s, 9.0, 1.55, '空间体验', 'Spatial Experience', tl, 700)
        tl.add(add_text(s, 9.0, 2.0, 3.75, 1.3, intro, size=10.5, color='ink_soft', font=F['body'], line=1.4), 'fade', at=900, dur=700)
        fills = ['primary', 'accent', 'line']
        for i, (ic, k, v) in enumerate(points[:3]):
            yy = 3.45 + i * 1.12
            for b in self.icon_badge(s, ic, 9.02, yy, d=0.5, ring=False, fill=fills[i]):
                tl.add(b, 'zoom', at=1200 + i * 220, dur=400)
            t = add_text(s, 9.72, yy - 0.04, 3.05, 1.0, [[dict(text=k, font=F['head'], bold=True, color='ink', size=12.5)], v], size=10, color='ink_soft', font=F['body'], line=1.3)
            tl.add(t, 'slide', at=1300 + i * 220, dur=600, dx=0.03)
        return self.finish(s, tl, note)

    def zones(self, title, en, img, meta, zone_of, zones, note=''):
        """zone_of: {room_label: zone_index} ; zones: [(cn, en, chip_hex, rooms_text, desc)] (pill colours = primary family)"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        m = json.load(open(self.A(meta), encoding='utf8'))
        x, y, w, h, sc = self.fit(img, (0.75, 1.35, 7.9, 5.95))
        tl.add(add_pic(s, self.A(img), x, y, w, h, cover=False), 'zoom', at=300, dur=1200, scale=0.93)
        pillc = ['accent', 'primary', 'line', 'accent2']
        self._labels_on(s, tl, m['labels'], list(zone_of), (x, y), sc, colors={k: pillc[v] for k, v in zone_of.items()}, at=1400)
        for i, (cn, ens, hx, rooms, txt) in enumerate(zones):
            yy = 1.6 + i * 1.72
            chip = add_shape(s, 'frame', 9.0, yy, 0.5, 0.75, fill=hx, line='line', lw=0.5)
            t1 = add_text(s, 9.7, yy - 0.05, 3.1, 0.4, [[dict(text=cn + '  ', bold=True), dict(text=ens, latin=True, italic=True, size=8.5, color='line', bold=False)]], size=15, color='primary_dk', font=F['head'])
            t2 = add_text(s, 9.7, yy + 0.36, 3.1, 0.3, rooms, size=10, color='accent', font=F['head'], bold=True)
            t3 = add_text(s, 9.7, yy + 0.7, 3.05, 0.85, txt, size=10, color='ink_soft', font=F['body'], line=1.35)
            at = 1400 + i * 450
            tl.add(chip, 'wipe', at=at, dur=500, dir='up'); tl.add(t1, 'slide', at=at + 100, dur=550, dx=0.03)
            tl.add(t2, 'fade', at=at + 200, dur=450); tl.add(t3, 'fade', at=at + 300, dur=500)
        return self.finish(s, tl, note)

    def flows(self, title, en, plan_img, flows, entry=None, summary=None, offset=(0, 0), note=''):
        """flows: [(name, color, dash|None, [path,...], desc)] path = [(px,py),...] source pixels ;
        entry: (px, py, label) ; summary: (title, text)"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        x, y, w, h, sc = self.fit(plan_img, (0.95, 1.45, 6.4, 5.65))
        size = Image.open(self.A(plan_img)).size
        self._plan_frame(s, x, y, w, h, tl)
        tl.add(add_pic(s, self.A(plan_img), x, y, w, h, cover=False, alpha=80), 'fade', at=300, dur=900)
        for i, (nm, col, dash, paths, _) in enumerate(flows):
            for j, pth in enumerate(paths):
                pts = [self.pmap(px, py, (x, y, w, h), size, offset) for px, py in pth]
                tl.add(freeform(s, pts, col, width=2.5, dash=dash), 'wipe', at=1200 + i * 700 + j * 200, dur=700)
        if entry:
            ex, ey = self.pmap(entry[0], entry[1], (x, y, w, h), size, offset)
            tl.add(add_shape(s, 'oval', ex - 0.17, ey - 0.17, 0.34, 0.34, fill='accent2', line='white', lw=1), 'zoom', at=1000, dur=450)
            tl.add(add_text(s, ex + 0.22, ey - 0.14, 0.7, 0.28, entry[2], size=10, color='accent2', font=F['head'], bold=True), 'fade', at=1100, dur=450)
        self.side_heading(s, 7.85, 1.5, '动线图例', 'Legend', tl, 700)
        for i, (nm, col, dash, _, txt) in enumerate(flows):
            yy = 2.0 + i * 1.2
            at = 1200 + i * 700
            tl.add(add_line(s, 7.9, yy + 0.18, 8.55, yy + 0.18, color=col, width=2.5, dash=dash), 'wipe', at=at, dur=500)
            tl.add(add_text(s, 8.7, yy, 4.0, 0.36, nm, size=13.5, color='ink', font=F['head'], bold=True), 'fade', at=at + 100, dur=450)
            tl.add(add_text(s, 8.7, yy + 0.38, 4.05, 0.75, txt, size=10, color='ink_soft', font=F['body'], line=1.35), 'fade', at=at + 200, dur=500)
        if summary:
            c = self.card(s, 7.85, 5.75, 4.9, 1.15, alpha=75)
            t = add_text(s, 8.05, 5.82, 4.55, 1.0, [[dict(text=summary[0], font=F['head'], bold=True, color='accent', size=12.5)], summary[1]],
                         size=10, color='ink_soft', font=F['body'], line=1.35, anchor='m')
            tl.add(c, 'wipe', at=3300, dur=600); tl.add(t, 'fade', at=3500, dur=700)
        return self.finish(s, tl, note)

    def _keyplan(self, s, x, y, w, plan_img, rect, offset, tl, at):
        size = Image.open(self.A(plan_img)).size
        h = w * size[1] / size[0]
        fr = add_shape(s, 'rect', x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill='white', alpha=60, line='line', lw=0.5)
        p = add_pic(s, self.A(plan_img), x, y, w, h, cover=False)
        x0, y0 = self.pmap(rect[0], rect[1], (x, y, w, h), size, offset)
        x1, y1 = self.pmap(rect[2], rect[3], (x, y, w, h), size, offset)
        hl = add_shape(s, 'rect', x0, y0, x1 - x0, y1 - y0, fill='accent', alpha=28, line='accent', lw=1.75)
        tl.add(fr, 'fade', at=at, dur=500); tl.add(p, 'fade', at=at, dur=600); tl.add(hl, 'zoom', at=at + 500, dur=500)
        return h

    def room(self, title, en, photo, points, keyplan=None, rect=None, offset=(0, 0), side='right', details=(), focus=(0.5, 0.5), note=''):
        """big render + key plan (room highlighted) + 3-4 icon points ; side='left' mirrors ; details: [crop paths] (arched)"""
        s, tl, n = self.new('gallery')
        self.chrome(s, n, tl, title, en, sprig=False)
        px, pw, rx, rw = (0.95, 7.85, 9.2, 3.55) if side == 'right' else (4.6, 7.95, 0.95, 3.25)
        py, ph = 1.48, 5.3
        o = add_shape(s, 'rect', px - 0.12, py - 0.12, pw + 0.24, ph + 0.24, line='line', lw=0.75)
        p = add_pic(s, self.A(photo), px, py, pw, ph, focus=focus)
        shadow(p, blur=0.3, dist=0.07, alpha=30)
        tl.add(p, 'zoom', at=200, dur=1200, scale=0.94); tl.add(o, 'fade', at=700, dur=800)
        ty = 1.5
        if keyplan:
            kw = 1.75 if details else 2.05
            kh = self._keyplan(s, rx, 1.5, kw, keyplan, rect, offset, tl, 900)
            if details:
                k = len(details)
                dw = min(0.85, (rw - kw - 0.18 - 0.1 * (k - 1)) / k)
                for i, ph_ in enumerate(details):
                    d = add_pic(s, self.A(ph_), rx + kw + 0.18 + i * (dw + 0.1), 1.5, dw, kh, geom='frame', focus=(0.5, 0.4))
                    pic_outline(d, 'line', 0.6)
                    tl.add(d, 'rise', at=1000 + i * 150, dur=700, dy=0.05)
            ty = 1.5 + kh + 0.25
        self.points_list(s, rx, ty, rw, points, tl, 1300, gap=(6.95 - ty) / max(len(points), 1))
        return self.finish(s, tl, note)

    def board(self, title, en, img, points, spec_lines, note=''):
        """exhibition board image (scripts/board) + composition notes"""
        s, tl, n = self.new('board')
        self.chrome(s, n, tl, title, en, sprig=False)
        if self.has(img):
            x, y, w, h, sc = self.fit(img, (0.95, 1.45, 8.6, 5.55))
            o = add_shape(s, 'rect', x - 0.1, y - 0.1, w + 0.2, h + 0.2, line='line', lw=0.75)
            p = add_pic(s, self.A(img), x, y, w, h, cover=False)
            shadow(p, blur=0.35, dist=0.09, alpha=32)
            tl.add(p, 'zoom', at=300, dur=1300, scale=0.9); tl.add(o, 'fade', at=900, dur=700)
        self.side_heading(s, 9.95, 1.55, '版面构成', 'Layout', tl, 800)
        self.points_list(s, 9.95, 2.1, 2.85, points, tl, 1100, gap=0.95, size=10)
        sp = add_text(s, 9.95, 6.0, 2.9, 0.9, [[dict(text=k + '　', font=F['head'], bold=True, color='primary_dk'), dict(text=v)] for k, v in spec_lines],
                      size=10, color='ink_soft', font=F['body'], line=1.4)
        tl.add(sp, 'fade', at=2100, dur=700)
        return self.finish(s, tl, note)

    def summary(self, title, en, quote, quote_en, tiles, note=''):
        """tiles: [(icon, key, text, photo)]x4"""
        s, tl, n = self.new()
        self.chrome(s, n, tl, title, en, sprig=False)
        q = add_text(s, 1.0, 1.55, 11.3, 0.6, quote, size=22, color='primary_dk', font=F['head'], bold=True, align='c', anchor='m')
        qe = add_text(s, 1.0, 2.15, 11.3, 0.3, quote_en, size=9.5, color='line', latin=F['latin'], italic=True, spacing=2.5, align='c')
        tl.add(q, 'letters', at=400, dur=400, pct=9000); tl.add(qe, 'wipe', at=1400, dur=700)
        fills = ['primary', 'accent', 'extra', 'line']
        for i, (ic, k, v, ph) in enumerate(tiles[:4]):
            x = 1.0 + i * 2.9
            c = self.card(s, x, 2.75, 2.6, 4.1, alpha=65, kind='frame')
            pic = add_pic(s, self.A(ph), x + 0.25, 2.97, 2.1, 1.75, geom='frame', focus=(0.5, 0.45))
            bd = self.icon_badge(s, ic, x + 1.02, 4.42, d=0.56, fill=fills[i])
            t1 = add_text(s, x + 0.1, 5.15, 2.4, 0.38, k, size=13.5, color='ink', font=F['head'], bold=True, align='c')
            t2 = add_text(s, x + 0.25, 5.58, 2.1, 1.1, v, size=10.5, color='ink_soft', font=F['body'], line=1.4, align='c')
            at = 1000 + i * 260
            tl.add(c, 'rise', at=at, dur=900, dy=0.05); tl.add(pic, 'fade', at=at + 200, dur=800)
            for b in bd:
                tl.add(b, 'zoom', at=at + 500, dur=400)
            tl.add(t1, 'fade', at=at + 600, dur=500); tl.add(t2, 'fade', at=at + 700, dur=600)
        return self.finish(s, tl, note)

    def thanks(self, photo, focus=(0.5, 0.5), big='THANKS', sub='感谢聆听', line='欢迎各位老师与同学批评指正', footer=None, note='谢谢聆听。'):
        """footer: list of run dicts or str"""
        s, tl, n = self.new('end')
        q = self.deco(s, 'ornament', 4.9, 0.55, w=3.5, alpha=22)
        p, fr = self.framed_photo(s, photo, 7.6, 0.95, 4.6, 5.95, frame=0.17, focus=focus, double=True)
        ob = self.deco(s, 'branch_tr', 9.5, -0.42, w=4.3)
        rv = self.deco(s, 'garland', 5.9, 5.75, w=3.2)
        lv = self.deco(s, 'sprig_br', -0.2, 5.5, h=2.3)
        t = add_text(s, 1.0, 1.75, 6.2, 1.5, big, size=80, color='primary_dk', latin=F['latin'], font=F['head'], spacing=4)
        t2 = add_text(s, 1.08, 3.3, 6, 0.6, sub, size=28, color='ink', font=F['head'], bold=True, spacing=6)
        dv = self.deco(s, 'divider', 1.05, 4.05, w=4.6)
        t3 = add_text(s, 1.08, 4.45, 6, 0.4, line, size=14, color='ink_soft', font=F['body'])
        t4 = add_text(s, 1.08, 5.0, 6, 0.4, [footer] if isinstance(footer, list) else (footer or ''), size=12, color='ink_soft', font=F['body'])
        seal = self.deco(s, 'seal', 3.55, 3.22, w=0.75)
        tl.add(q, 'fade', at=0, dur=1200); tl.add(p, 'zoom', at=0, dur=1300, scale=0.9)
        for i, f in enumerate(fr):
            tl.add(f, 'fade', at=500 + i * 200, dur=900)
        tl.add(ob, 'rise', at=800, dur=1200, dy=-0.06); tl.add(rv, 'rise', at=1000, dur=1200, dy=0.06); tl.add(lv, 'rise', at=1100, dur=1100, dy=0.06)
        tl.add(t, 'letters', at=400, dur=500, pct=15000); tl.add(t2, 'rise', at=1200, dur=900); tl.add(dv, 'split', at=1700, dur=600)
        tl.add(t3, 'fade', at=1900, dur=700); tl.add(t4, 'fade', at=2100, dur=700); tl.add(seal, 'zoom', at=2300, dur=450)
        return self.finish(s, tl, note)
